"""XLSX export with provenance.

Exports a Result to an Excel workbook with three sheets:
  - Comparison: the financial table with hyperlinked numeric cells + comments
  - Provenance: one row per distinct fact (raw + YTD_DIFF derived + ratio derived)
  - Claims: claim text, evidence, verdict, and reason

Also writes a JSON snapshot of all facts alongside the workbook.
"""
import json
import pathlib
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.comments import Comment

from finagent.datasource.facts import quarter_facts
from finagent.metrics import reconstruct_ratio_derived_facts


def _build_fact_index(result):
    """Build a comprehensive fact_id -> entry lookup covering all three sources:

    1. Native raw facts (type "raw")
    2. YTD-differenced single-quarter facts from quarter_facts() (type "ytd_diff")
    3. Ratio-derived facts reconstructed from table rows with formula (type "ratio")

    Returns: dict[fact_id -> dict] where each entry has:
      - "type": "raw" | "ytd_diff" | "ratio"
      - "fact": Fact object (for raw/ytd_diff)
      - For ratio: "numerator", "denominator" (Fact objects), "derived" (DerivedFact)
    """
    index = {}

    for company, raw_facts in result.facts_by_company.items():
        # Source 1 + 2: raw facts AND quarter_facts output (covers native + YTD_DIFF)
        all_qf = quarter_facts(raw_facts)

        for f in all_qf:
            if f.provenance.raw_value == "derived:YTD_DIFF":
                index[f.fact_id] = {"type": "ytd_diff", "fact": f}
            else:
                index[f.fact_id] = {"type": "raw", "fact": f}

        # Source 3: ratio formula rows — reconstruct DerivedFact per cell frame
        # Uses the shared helper from metrics.py (same one used by app.py for
        # the evidence bundle) to ensure consistent fact IDs and values.
        ratio_derived = reconstruct_ratio_derived_facts(raw_facts, result.table.rows, company.ticker)
        # Build a fact_id -> Fact lookup for input facts so we can attach
        # numerator/denominator metadata for comments and hyperlinks.
        fact_by_id = {f.fact_id: f for f in all_qf}
        for d in ratio_derived:
            gp = fact_by_id.get(d.input_fact_ids[0])
            rv = fact_by_id.get(d.input_fact_ids[1])
            if gp is None or rv is None:
                continue
            index[d.fact_id] = {
                "type": "ratio",
                "numerator": gp,
                "denominator": rv,
                "derived": d,
            }

    return index


def _fact_comment_text(entry):
    """Build comment text for a cell from its fact index entry."""
    t = entry["type"]
    if t in ("raw", "ytd_diff"):
        f = entry["fact"]
        return (
            f"{f.concept}\n"
            f"{f.provenance.accession}\n"
            f"filed {f.provenance.filed_date}\n"
            f"{f.provenance.raw_value} {f.unit}"
        )
    elif t == "ratio":
        gp = entry["numerator"]
        rv = entry["denominator"]
        d = entry["derived"]
        pct = (d.value * 100).quantize(Decimal("0.000001"))
        return (
            f"derived:gross_profit/revenue\n"
            f"numerator: {gp.fact_id} = {gp.value} {gp.unit}\n"
            f"denominator: {rv.fact_id} = {rv.value} {rv.unit}\n"
            f"= {pct}%\n"
            f"source: {gp.provenance.accession}\n"
            f"filed {gp.provenance.filed_date}"
        )
    return ""


def _fact_hyperlink(entry):
    """Return the source_url to hyperlink to for a fact index entry."""
    t = entry["type"]
    if t in ("raw", "ytd_diff"):
        return entry["fact"].provenance.source_url
    elif t == "ratio":
        # Ratio cells hyperlink to the numerator (gross profit) fact's source_url
        return entry["numerator"].provenance.source_url
    return None


def _build_snapshot_facts(fact_index):
    """Convert fact_index entries into the snapshot JSON list format."""
    out = []
    for fid, entry in fact_index.items():
        t = entry["type"]
        if t in ("raw", "ytd_diff"):
            f = entry["fact"]
            out.append({
                "fact_id": f.fact_id,
                "ticker": f.ticker,
                "company": f.company_name,
                "concept": f.concept,
                "period_start": str(f.period_start),
                "period_end": str(f.period_end),
                "frame": f.frame,
                "raw_value": f.provenance.raw_value,
                "value": str(f.value),
                "unit": f.unit,
                "form": f.provenance.form,
                "accession": f.provenance.accession,
                "filed_date": f.provenance.filed_date,
                "source_url": f.provenance.source_url,
                "type": t,
            })
        elif t == "ratio":
            d = entry["derived"]
            gp = entry["numerator"]
            out.append({
                "fact_id": d.fact_id,
                "ticker": gp.ticker,
                "company": gp.company_name,
                "concept": "gross_profit/revenue",
                "period_start": str(gp.period_start),
                "period_end": str(gp.period_end),
                "frame": gp.frame,
                "raw_value": None,
                "value": str(d.value),
                "unit": gp.unit,
                "form": gp.provenance.form,
                "accession": gp.provenance.accession,
                "filed_date": gp.provenance.filed_date,
                "source_url": gp.provenance.source_url,
                "type": "ratio",
                "formula": "gross_profit/revenue",
                "input_fact_ids": list(d.input_fact_ids),
            })
    return out


def export_xlsx(result, path):
    """Export a Result to an XLSX workbook + JSON snapshot.

    Args:
        result: a Result object (plan, table, claims, verdicts, facts_by_company, ...)
        path: output path for the .xlsx file

    Returns:
        pathlib.Path of the written xlsx file

    Side effect:
        Writes a JSON snapshot to <parent>/snapshots/<stem>.json
    """
    path = pathlib.Path(path)

    # Build the comprehensive fact index (raw + ytd_diff + ratio)
    fact_index = _build_fact_index(result)

    # Compute snapshot facts once — used by both Provenance sheet and JSON snapshot
    snapshot_facts = _build_snapshot_facts(fact_index)

    # --- Snapshot JSON (written first so a save failure leaves no orphan workbook) ---
    snap_dir = path.parent / "snapshots"
    snap_dir.mkdir(parents=True, exist_ok=True)
    snap_path = snap_dir / (path.stem + ".json")

    snapshot = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "facts": snapshot_facts,
    }
    snap_path.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False))

    # --- Workbook ---
    wb = Workbook()

    # --- Comparison sheet ---
    cmp = wb.active
    cmp.title = "Comparison"

    # Header: "Company / Metric" + column frames
    cmp.append(["Company / Metric", *result.table.columns])

    for row in result.table.rows:
        row_label = f"{row.company} {row.alias}"
        cmp.append([row_label] + [
            (str(c.value) if c.value is not None else c.note) for c in row.cells
        ])
        r = cmp.max_row
        for i, c in enumerate(row.cells, start=2):
            if c.fact_id and c.fact_id in fact_index:
                entry = fact_index[c.fact_id]
                url = _fact_hyperlink(entry)
                if url:
                    cmp.cell(r, i).hyperlink = url
                comment_text = _fact_comment_text(entry)
                if comment_text:
                    cmp.cell(r, i).comment = Comment(comment_text, "finagent")

    # --- Provenance sheet (reuses snapshot facts — computed once) ---
    prov = wb.create_sheet("Provenance")
    prov.append([
        "fact_id", "ticker", "company", "concept", "period_start", "period_end",
        "frame", "raw_value", "value", "unit", "form", "accession",
        "filed_date", "source_url", "type",
    ])
    for entry in snapshot_facts:
        prov.append([
            entry["fact_id"],
            entry["ticker"],
            entry["company"],
            entry["concept"],
            entry["period_start"],
            entry["period_end"],
            entry["frame"],
            entry["raw_value"],
            entry["value"],
            entry["unit"],
            entry["form"],
            entry["accession"],
            entry["filed_date"],
            entry["source_url"],
            entry["type"],
        ])

    # --- Claims sheet ---
    cl = wb.create_sheet("Claims")
    cl.append(["claim_id", "text", "type", "evidence_ids", "verdict", "reason"])
    reasons = {v.claim_id: v for v in result.verdicts}
    for c in result.claims:
        v = reasons.get(c.id)
        cl.append([
            c.id,
            c.text,
            c.claim_type,
            ",".join(c.evidence_ids),
            "pass" if v and v.ok else "fail",
            v.reason if v else "",
        ])

    # Save workbook (after snapshot — a save failure leaves only the harmless snapshot)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    return path
